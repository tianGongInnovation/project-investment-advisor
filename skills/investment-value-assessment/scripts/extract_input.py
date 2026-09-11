# -*- coding: utf-8 -*-
"""提取输入文档（docx/pptx/pdf/txt）全文文本，供投资价值评估使用。

用法：
    python extract_input.py <文件或目录路径> [输出目录]

- 支持递归处理目录。
- 输出为同名 .txt（UTF-8），默认输出到 源文件同目录 的 _extracted 子目录。
- pdf 优先用 pypdf，失败时提示改用 markitdown。
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")


def extract_docx(path):
    from docx import Document
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n===== 幻灯片 {i} =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text_frame.text.strip().replace("\n", " ") for c in row.cells]
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            raise RuntimeError("缺少 pypdf，请先安装：pip install pypdf（或改用 markitdown 转换）")
    reader = PdfReader(path)
    lines = []
    for i, page in enumerate(reader.pages, 1):
        lines.append(f"\n===== 第 {i} 页 =====")
        lines.append(page.extract_text() or "(本页无可提取文本，可能是扫描件，需OCR)")
    return "\n".join(lines)


def extract_txt(path):
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise RuntimeError("无法识别文本编码")


EXTRACTORS = {
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".pdf": extract_pdf,
    ".txt": extract_txt,
    ".md": extract_txt,
}


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    target = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else None
    files = []
    if os.path.isdir(target):
        for root, dirs, names in os.walk(target):
            dirs[:] = [d for d in dirs if d != "_extracted"]
            for n in names:
                files.append(os.path.join(root, n))
    else:
        files.append(target)

    for path in files:
        ext = os.path.splitext(path)[1].lower()
        if ext not in EXTRACTORS:
            continue
        try:
            text = EXTRACTORS[ext](path)
        except Exception as e:
            print(f"ERR {os.path.basename(path)}: {e}")
            continue
        base = os.path.splitext(os.path.basename(path))[0]
        dest_dir = out_dir or os.path.join(os.path.dirname(path), "_extracted")
        os.makedirs(dest_dir, exist_ok=True)
        outp = os.path.join(dest_dir, base + ".txt")
        with open(outp, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"OK  {path} -> {outp} ({len(text)} 字符)")


if __name__ == "__main__":
    main()
